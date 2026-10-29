"""Exportación de la síntesis a un informe Word con estilo profesional."""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path


def save_summary_docx(
    text: str,
    path: str | Path,
    *,
    source_name: str = "Texto pegado",
    engine: str = "Motor local",
) -> None:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:
        raise ImportError(
            "Falta python-docx. Instálalo con: pip install -r requirements.txt"
        ) from exc

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(45, 55, 72)
    normal.paragraph_format.space_after = Pt(5)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("Síntesis ejecutiva")
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(23)
    run.font.color.rgb = RGBColor(27, 55, 87)

    meta = document.add_paragraph()
    meta.paragraph_format.space_after = Pt(14)
    meta_run = meta.add_run(
        f"Documento: {Path(source_name).name}    ·    "
        f"Generado: {date.today().strftime('%d/%m/%Y')}    ·    {engine}"
    )
    meta_run.font.size = Pt(8.5)
    meta_run.font.color.rgb = RGBColor(100, 116, 139)

    section_names = {
        "Resumen",
        "Ideas / conceptos clave",
        "Datos, cifras, fechas o definiciones",
        "Decisiones",
        "Acciones propuestas o pendientes",
        "Riesgos y dudas",
        "Estructura / apartados",
        "Conclusiones o implicaciones",
        "Identificación",
        "Especificaciones",
        "Conexión / instalación",
        "Programas",
        "Uso diario",
        "Mantenimiento",
        "Códigos de error / averías",
        "Advertencias importantes",
        "Matriz Resumen de Indicadores",
    }
    started = False

    def add_heading(label: str, *, size) -> None:
        paragraph = document.add_paragraph(style="Heading 2")
        paragraph.paragraph_format.keep_with_next = True
        heading_run = paragraph.add_run(label)
        heading_run.font.name = "Aptos Display"
        heading_run.font.size = size
        heading_run.font.color.rgb = RGBColor(31, 77, 58)

    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("──") and line.endswith("──"):
            line = line.strip("─ ").strip()
            if not line:
                continue
        numbered = bool(
            re.match(r"^\d+\.\s+\S", line) and len(line) <= 80 and not line.endswith(".")
        )
        if line in section_names or numbered:
            add_heading(line, size=Pt(12.5))
            started = True
            continue
        if not started:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(8)
            title_run = paragraph.add_run(line)
            title_run.bold = True
            title_run.font.name = "Aptos Display"
            title_run.font.size = Pt(16)
            title_run.font.color.rgb = RGBColor(27, 55, 87)
            started = True
            continue
        if line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            content = line[2:]
            if ": " in content:
                label, value = content.split(": ", 1)
                lead = paragraph.add_run(f"{label}: ")
                lead.bold = True
                paragraph.add_run(value)
            else:
                paragraph.add_run(content)
            continue
        if ": " in line and len(line.split(":", 1)[0].split()) <= 8:
            paragraph = document.add_paragraph()
            label, value = line.split(": ", 1)
            lead = paragraph.add_run(f"{label}: ")
            lead.bold = True
            paragraph.add_run(value)
            continue
        document.add_paragraph(line)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run(
        "FichaManual · Síntesis automática; revise los datos antes de distribuir"
    )
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(120, 130, 140)

    properties = document.core_properties
    properties.title = "Síntesis ejecutiva — FichaManual"
    properties.subject = f"Resumen de {Path(source_name).name}"
    properties.author = "FichaManual"
    document.save(str(path))
